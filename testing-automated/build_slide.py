"""Generate a single-slide PowerPoint summarising the automated testing harness."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

OUT = Path(__file__).parent / "automated-testing-overview.pptx"

NAVY = RGBColor(0x03, 0x2D, 0x60)
BLUE = RGBColor(0x00, 0xA1, 0xE0)
LIGHT = RGBColor(0xF4, 0xF6, 0xF9)
GREY = RGBColor(0x54, 0x69, 0x8D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x2E, 0x84, 0x4A)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

# Background band
band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.1))
band.fill.solid(); band.fill.fore_color.rgb = NAVY
band.line.fill.background()

# Title
title_box = slide.shapes.add_textbox(Inches(0.4), Inches(0.18), Inches(12.5), Inches(0.8))
tf = title_box.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Scalable, Repeatable, Automated Agentforce Testing"
p.runs[0].font.size = Pt(28); p.runs[0].font.bold = True
p.runs[0].font.color.rgb = WHITE

# Subtitle
sub = slide.shapes.add_textbox(Inches(0.4), Inches(0.65), Inches(12.5), Inches(0.4))
sp = sub.text_frame.paragraphs[0]
sp.text = "LNA Nova SDR Agent · run-automated-tests.py"
sp.runs[0].font.size = Pt(14); sp.runs[0].font.color.rgb = WHITE
sp.runs[0].font.italic = True

# ---------- Pipeline (4 steps) ----------
steps = [
    ("1. Seed Data", "Create N test Leads in the org from JSON;\npersist 15-char Salesforce Ids."),
    ("2. Generate", "Render YAML test suites from templates,\none case per Lead, with dynamic context."),
    ("3. Deploy & Run", "sf agent test create / run / results\non Agentforce_Sales_Development_Rep."),
    ("4. Report", "Parse outcomes, capture drafted emails,\nwrite timestamped Markdown report."),
]

step_top = Inches(1.4)
step_h = Inches(1.4)
step_w = Inches(2.85)
gap = Inches(0.25)
left = Inches(0.4)

for i, (head, body) in enumerate(steps):
    x = left + (step_w + gap) * i
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, step_top, step_w, step_h)
    box.fill.solid(); box.fill.fore_color.rgb = LIGHT
    box.line.color.rgb = BLUE; box.line.width = Pt(1.25)
    tf = box.text_frame; tf.margin_left = Inches(0.12); tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.1); tf.word_wrap = True

    p1 = tf.paragraphs[0]
    p1.text = head
    p1.runs[0].font.size = Pt(15); p1.runs[0].font.bold = True
    p1.runs[0].font.color.rgb = NAVY

    p2 = tf.add_paragraph()
    p2.text = body
    p2.runs[0].font.size = Pt(11); p2.runs[0].font.color.rgb = GREY

    # Arrow between steps
    if i < len(steps) - 1:
        arr_x = x + step_w + Inches(0.02)
        arr = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, arr_x, step_top + Inches(0.55), Inches(0.21), Inches(0.3))
        arr.fill.solid(); arr.fill.fore_color.rgb = BLUE
        arr.line.fill.background()

# ---------- Three pillars ----------
pillars = [
    ("Scalable",
     "One template -> N leads x M suites.\n"
     "Add a lead in JSON or a new template;\nno code change required."),
    ("Repeatable",
     "Single timestamped run-id stamps Leads,\n"
     "suite api-names, output folders.\n"
     "Re-runs never collide."),
    ("Automated",
     "End-to-end via sf CLI: seed -> generate ->\n"
     "deploy -> execute -> parse -> report.\n"
     "--step / --skip-leads for partial runs."),
]

pil_top = Inches(3.05)
pil_h = Inches(1.85)
pil_w = Inches(4.05)
pil_gap = Inches(0.18)
pil_left = Inches(0.4)

for i, (head, body) in enumerate(pillars):
    x = pil_left + (pil_w + pil_gap) * i
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, pil_top, pil_w, pil_h)
    card.fill.solid(); card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = NAVY; card.line.width = Pt(1.5)
    tf = card.text_frame; tf.margin_left = Inches(0.18); tf.margin_top = Inches(0.15)
    tf.margin_right = Inches(0.18); tf.word_wrap = True

    # Header bar
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, pil_top, pil_w, Inches(0.05))
    bar.fill.solid(); bar.fill.fore_color.rgb = BLUE
    bar.line.fill.background()

    p1 = tf.paragraphs[0]
    p1.text = head
    p1.runs[0].font.size = Pt(20); p1.runs[0].font.bold = True
    p1.runs[0].font.color.rgb = NAVY

    p2 = tf.add_paragraph()
    p2.text = body
    p2.runs[0].font.size = Pt(12); p2.runs[0].font.color.rgb = GREY
    p2.space_before = Pt(6)

# ---------- Footer metrics strip ----------
footer_top = Inches(5.15)
footer_h = Inches(0.85)
footer = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), footer_top, Inches(12.55), footer_h)
footer.fill.solid(); footer.fill.fore_color.rgb = NAVY
footer.line.fill.background()

metrics = [
    ("4", "Test suites covered"),
    ("6+", "Personas tested in parallel"),
    ("3", "CLI flags for partial runs"),
    ("100%", "Outputs versioned by timestamp"),
]
m_w = Inches(12.55) / len(metrics)
for i, (num, label) in enumerate(metrics):
    mx = Inches(0.4) + m_w * i
    tb = slide.shapes.add_textbox(mx, footer_top + Inches(0.05), m_w, footer_h)
    tf = tb.text_frame; tf.word_wrap = True
    p1 = tf.paragraphs[0]; p1.alignment = PP_ALIGN.CENTER
    p1.text = num
    p1.runs[0].font.size = Pt(22); p1.runs[0].font.bold = True
    p1.runs[0].font.color.rgb = WHITE
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    p2.text = label
    p2.runs[0].font.size = Pt(11); p2.runs[0].font.color.rgb = LIGHT

# ---------- Tagline ----------
tag = slide.shapes.add_textbox(Inches(0.4), Inches(6.15), Inches(12.55), Inches(1.15))
tf = tag.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
p.text = "From hours of manual UAT to a single command — every release, regression-tested."
p.runs[0].font.size = Pt(16); p.runs[0].font.italic = True
p.runs[0].font.color.rgb = NAVY

p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
p2.text = "python3 run-automated-tests.py --org <alias>"
p2.runs[0].font.size = Pt(13); p2.runs[0].font.bold = True
p2.runs[0].font.name = "Courier New"
p2.runs[0].font.color.rgb = GREEN
p2.space_before = Pt(8)

prs.save(OUT)
print(f"Wrote {OUT}")
